from __future__ import annotations

import json
import sys
from pathlib import Path

import fitz

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


def inspect_pdf(path: Path) -> dict:
    with fitz.open(path) as document:
        outline = document.get_toc()
        link_count = 0
        link_pages = 0
        bad_goto_links: list[dict] = []
        toc_like_pages: list[int] = []

        for page_index in range(document.page_count):
            page = document[page_index]
            text = page.get_text("text")
            if "TABLE OF CONTENTS" in text.upper() or "TABLE OF CONTENT" in text.upper():
                toc_like_pages.append(page_index + 1)

            links = page.get_links()
            if links:
                link_pages += 1
                link_count += len(links)

            for link in links:
                if link.get("kind") == fitz.LINK_GOTO:
                    target = link.get("page")
                    if target is None or target < 0 or target >= document.page_count:
                        bad_goto_links.append(
                            {
                                "source_page": page_index + 1,
                                "target_page": target,
                            }
                        )

        return {
            "path": str(path),
            "exists": True,
            "pages": document.page_count,
            "outline_count": len(outline),
            "outline_sample": outline[:5],
            "toc_like_pages": toc_like_pages[:25],
            "toc_like_page_count": len(toc_like_pages),
            "link_pages": link_pages,
            "link_count": link_count,
            "bad_goto_link_count": len(bad_goto_links),
            "bad_goto_links_sample": bad_goto_links[:10],
        }


def inspect_json(path: Path) -> dict:
    if not path.exists():
        return {"path": str(path), "exists": False}
    data = json.loads(path.read_text(encoding="utf-8"))
    return {
        "path": str(path),
        "exists": True,
        "heading_count": len(data) if isinstance(data, list) else None,
        "sample": data[:3] if isinstance(data, list) else data,
    }


def inspect_xlsx(path: Path) -> dict:
    if load_workbook is None:
        return {
            "path": str(path),
            "exists": True,
            "error": "openpyxl is not installed in this Python environment",
        }
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    headers = [str(value) if value is not None else "" for value in (rows[0] if rows else [])]
    records = []
    for row in rows[1:6]:
        records.append({headers[index]: value for index, value in enumerate(row) if index < len(headers)})
    return {
        "path": str(path),
        "exists": True,
        "sheet": worksheet.title,
        "data_rows": max(0, len(rows) - 1),
        "headers": headers,
        "sample": records,
    }


def main() -> int:
    for raw_path in sys.argv[1:]:
        path = Path(raw_path)
        if not path.exists():
            print(json.dumps({"path": str(path), "exists": False}, indent=2))
        elif path.suffix.lower() == ".pdf":
            print(json.dumps(inspect_pdf(path), indent=2))
        elif path.suffix.lower() == ".json":
            print(json.dumps(inspect_json(path), indent=2))
        elif path.suffix.lower() == ".xlsx":
            print(json.dumps(inspect_xlsx(path), indent=2, default=str))
        else:
            print(json.dumps({"path": str(path), "exists": True}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
