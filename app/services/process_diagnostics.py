from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.models import Heading
from app.services.existing_toc_linker import ExistingTocLinkResult, ExistingTocRow

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def write_process_report(
    *,
    document_id: str,
    output_dir: Path,
    mode: str,
    source_filename: str,
    summary: dict[str, Any],
    unresolved_rows: list[ExistingTocRow] | None = None,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = output_dir / f"{document_id}_process.xlsx"

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    _build_summary_sheet(
        summary_sheet,
        document_id=document_id,
        source_filename=source_filename,
        mode=mode,
        summary=summary,
    )

    detail_sheet = workbook.create_sheet("Unresolved Links")
    _build_unresolved_sheet(detail_sheet, unresolved_rows or [])

    workbook.save(report_path)
    _write_preview_text(report_path, summary, unresolved_rows or [])
    return report_path


def summarize_heading_stats(headings: list[Heading]) -> dict[str, Any]:
    if not headings:
        return {
            "headings_extracted": 0,
            "pages_with_headings": 0,
            "sparse_heading_pages": "",
        }
    per_page: dict[int, int] = {}
    for heading in headings:
        per_page[heading.page] = per_page.get(heading.page, 0) + 1
    sparse_pages = sorted(page for page, count in per_page.items() if count == 1)
    return {
        "headings_extracted": len(headings),
        "pages_with_headings": len(per_page),
        "sparse_heading_pages": ", ".join(str(page) for page in sparse_pages[:25]),
    }


def summarize_link_stats(result: ExistingTocLinkResult) -> dict[str, Any]:
    return {
        "linked_rows": len(result.linked_rows),
        "unresolved_rows": len(result.unresolved_rows),
        "toc_pages_count": len(result.toc_pages),
        "reference_pages_count": len(result.reference_pages),
    }


def read_process_report_preview(report_path: Path) -> str:
    preview_path = report_path.with_suffix(".preview.txt")
    if preview_path.exists():
        return preview_path.read_text(encoding="utf-8")
    if not report_path.exists():
        return ""
    try:
        workbook = load_workbook(report_path, read_only=True, data_only=True)
        summary_sheet = workbook["Summary"] if "Summary" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        rows = []
        for row in summary_sheet.iter_rows(min_row=1, max_row=30, min_col=1, max_col=2, values_only=True):
            key, value = row
            if key is None and value is None:
                continue
            rows.append(f"{key}: {value}")
        return "\n".join(rows)
    except Exception:
        return "Process report available for download."


def _build_summary_sheet(
    sheet,
    *,
    document_id: str,
    source_filename: str,
    mode: str,
    summary: dict[str, Any],
) -> None:
    now = datetime.now(timezone.utc).isoformat()
    rows: list[tuple[str, Any]] = [
        ("timestamp_utc", now),
        ("document_id", document_id),
        ("source_filename", source_filename),
        ("mode", mode),
    ]
    rows.extend((key, value) for key, value in summary.items())

    sheet.append(["Field", "Value"])
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for key, value in rows:
        sheet.append([key, value])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:B{sheet.max_row}"
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 80
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")


def _build_unresolved_sheet(sheet, unresolved_rows: list[ExistingTocRow]) -> None:
    headers = [
        "toc_page",
        "toc_type",
        "target_label",
        "unresolved_reason",
        "title",
        "reference_text",
        "chapter",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for row in unresolved_rows:
        sheet.append(
            [
                row.page_number,
                row.toc_type,
                row.target_label or "",
                row.unresolved_reason or "target_not_found",
                row.title,
                row.reference_text,
                row.chapter,
            ]
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{max(sheet.max_row, 1)}"
    widths = {"A": 10, "B": 12, "C": 18, "D": 34, "E": 70, "F": 38, "G": 10}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=7):
        row[4].alignment = Alignment(wrap_text=True, vertical="top")
        row[5].alignment = Alignment(wrap_text=True, vertical="top")


def _write_preview_text(report_path: Path, summary: dict[str, Any], unresolved_rows: list[ExistingTocRow]) -> None:
    lines: list[str] = []
    for key, value in summary.items():
        lines.append(f"{key}: {value}")
    lines.append("")
    lines.append("Unresolved rows (first 20):")
    if not unresolved_rows:
        lines.append("None")
    else:
        for row in unresolved_rows[:20]:
            lines.append(
                f"- toc_page={row.page_number}, toc_type={row.toc_type}, "
                f"label={row.target_label or '-'}, reason={row.unresolved_reason or 'target_not_found'}"
            )
    report_path.with_suffix(".preview.txt").write_text("\n".join(lines), encoding="utf-8")
